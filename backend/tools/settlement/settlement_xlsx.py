#!/usr/bin/env python3
"""월 정산 엑셀 생성 — studio 정산 쿼리 TSV -> 정산_YYYYMM.xlsx.

사용: python3 settlement_xlsx.py <정산.tsv> <출력.xlsx> [월 라벨, 기본 '9월']
입력: 정산 쿼리 출력(탭 구분, 헤더 1줄, 마지막 '# rows=' 줄 무시)
시트: 안내 / 단가표 / 전체 작업 내역 / 방송사별 / 작업자별 요약 / 작업자별 시트
금액은 전부 수식 — 단가표 노란 칸만 채우면 계산됨.
"""
import csv, re, sys
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BRD = {"딜라이브": "DLIV", "KBS": "KBS", "LG헬로비전": "LGHV", "SK브로드밴드": "SKBB",
       "티캐스트": "TVCS", "TVCS": "TVCS", "TVING": "TVNG", "티빙": "TVNG"}
ORDER = ["DLIV", "KBS", "LGHV", "SKBB", "TVCS", "TVNG"]
WORK = {"SYNC": "싱크", "EDIT": "교정"}
# 시스템 등록 장르(t_prg.gnr_cd)가 틀린 프로그램 보정 — 안내 시트에 표시됨
GENRE_FIX = {"기쁜 우리 좋은날": "드라마"}

HEAD_FILL = PatternFill("solid", fgColor="44546A")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WON = '#,##0;-#,##0;"-"'


def episode_label(season, episode):
    if not episode:
        return None
    return f"시즌{season} {episode}회" if season else f"{episode}회"


def _num(v):
    return int(v) if re.fullmatch(r"\d+", v or "") else 10**9


def load(path):
    rows, skipped = [], []
    with open(path, encoding="utf-8-sig") as f:
        lines = [l for l in f if l.strip() and not l.startswith("#")]
    for r in csv.DictReader(lines, delimiter="\t" if "\t" in lines[0] else "|"):
        code = BRD.get(r["brd"].strip())
        if code is None:
            skipped.append(r); continue
        notes = []
        if r["done_by"] and r["done_by"] != r["assignee"]:
            notes.append(f"완료처리자: {r['done_by']}")
        if not r["dur_min_ceil"]:
            notes.append("영상길이 없음(t_vdo 미등록)")
        if r["work"] == "SYNC" and r["first_done_kst"] < "2026-09":
            notes.append("싱크 전월 완료 — 교정 9월 완료로 포함")
        rows.append({
            "brd": code, "genre": GENRE_FIX.get(r["program"], r["genre"]) or None, "program": r["program"],
            "season": r["season"], "episode": r["episode"],
            "ep": episode_label(r["season"], r["episode"]),
            "work": WORK[r["work"]], "assignee": r["assignee"] or None,
            "dur": int(r["dur_min_ceil"]) if r["dur_min_ceil"] else None,
            "work_min": float(r["work_min"]) if r["work_min"] else None,
            "done": r["first_done_kst"], "prj_id": int(r["prj_id"]),
            "warn": any(n.startswith(("완료처리자", "영상길이")) for n in notes),
            "note": " / ".join(notes) or None,
        })
    rows.sort(key=lambda x: (ORDER.index(x["brd"]), x["program"], _num(x["season"]),
                             _num(x["episode"]), x["prj_id"], x["work"] != "싱크"))
    return rows, skipped


def header(ws, row, titles, widths=None):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row, i, t)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), HEAD_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[chr(64 + i)].width = w


def put(ws, row, values, fmts=None, fill=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row, i, v)
        c.border = BOX
        if fmts and fmts[i - 1]:
            c.number_format = fmts[i - 1]
        if fill:
            c.fill = fill


def build(rows, skipped, out, month):
    wb = Workbook()
    n = len(rows); last = n + 1
    A = f"'전체 작업 내역'!$A$2:$A${last}"
    E = f"'전체 작업 내역'!$E$2:$E${last}"
    F = f"'전체 작업 내역'!$F$2:$F${last}"
    col = lambda L: f"'전체 작업 내역'!${L}$2:${L}${last}"

    # 안내
    ws = wb.active; ws.title = "안내"; ws.column_dimensions["A"].width = 140
    by_brd = defaultdict(int)
    for r in rows:
        by_brd[r["brd"]] += 1
    prev_sync = sum(1 for r in rows if r["work"] == "싱크" and r["done"] < "2026-09")
    no_dur = sum(1 for r in rows if r["dur"] is None)
    guide = [
        (f"{month} 작업자별 정산 데이터 (2026-09-01 ~ 2026-09-30)", True), (None, False),
        ("[입력 방법]", True),
        ("1. '단가표' 시트의 노란 칸(방송사 × 작업구분, 분당 단가)에 단가를 입력하면 모든 시트의 금액이 자동 계산됩니다.", False),
        ("2. 특정 회차만 단가가 다르면 '전체 작업 내역' G열(단가)에 숫자를 직접 덮어쓰면 됩니다. 작업자별 시트도 그 값을 따라갑니다.", False),
        ("3. 정산 금액 = 단가 × 영상 길이(분, 소수점 올림).", False), (None, False),
        ("[집계 기준]", True),
        ("· 대상: 교정 '최초 완료' 시각이 9/1 00:00 ~ 9/30 23:59(KST)인 회차. 그 회차의 싱크는 완료일과 관계없이 함께 포함(회차 전체를 교정 완료월에 정산).", False),
        (f"  - 싱크가 8월에 끝나고 교정이 9월에 끝난 회차: 싱크도 9월 정산에 포함 ({prev_sync}건, 비고 표시)", False),
        ("  - 싱크만 9월에 끝나고 교정이 10월 이후 완료(또는 미완료)인 회차: 9월에서 제외 → 교정 완료월에 정산", False),
        ("· 재작업(재싱크·재교정)은 중복 집계하지 않음.", False),
        ("· 담당자: 시스템에 배정된 역할별 담당자(싱크 담당 / 교정 담당).", False),
        ("· 영상 길이: 원본 영상 길이를 분 단위로 소수점 올림.", False),
        ("· 작업소요(분): 해당 작업 단계(싱크=싱크+재싱크, 교정=교정+재교정)에 편집기에서 실제 작업한 시간 합계. 참고용.", False),
        ("· 싱크 없이 교정만 있는 회차는 교정 1건만 표시됩니다.", False), (None, False),
        ("[확인 필요]", True),
        (f"· 주황색 행: 비고 확인 필요 — 영상 길이 없음({no_dur}건, 금액 0으로 계산됨) 또는 완료 처리자가 담당자와 다름.", False),
    ]
    for b in ORDER:
        if not by_brd[b]:
            guide.append((f"· {b}: 9월 완료 건 없음.", False))
    for s in skipped:
        guide.append((f"· 제외: 대상 6개 방송사 외 — {s['brd']} {s['program']} {s['episode']}회 "
                       f"{WORK.get(s['work'], s['work'])}(prj {s['prj_id']})", False))
    for prg, g in GENRE_FIX.items():
        guide.append((f"· 장르 보정: {prg} — 시스템 프로그램 장르는 예능으로 등록돼 있으나 회차 장르(수급)가 드라마라 {g}로 표기.", False))
    guide.append(("· 프로그램명은 시스템 등록명 그대로입니다(같은 프로그램이 다른 이름으로 등록된 경우 있음, 예: 태군노래자랑 / 태군 노래자랑 경북 경산).", False))
    for i, (t, bold) in enumerate(guide, 1):
        ws.cell(i, 1, t).font = Font(bold=bold, size=14 if i == 1 else 11)

    # 단가표
    ws = wb.create_sheet("단가표")
    ws["A1"] = "단가표 (분당 단가, 원)"; ws["A1"].font = Font(bold=True, size=13)
    header(ws, 3, ["방송사", "작업구분", "단가(원/분)"], [12, 12, 14])
    r = 4
    for b in ORDER:
        for w in ("싱크", "교정"):
            put(ws, r, [b, w, None], [None, None, "#,##0"])
            ws.cell(r, 3).fill = INPUT_FILL
            r += 1
    ws.cell(r + 1, 1, "노란 칸에 입력. 비워두면 단가 0으로 계산됩니다.")

    # 전체 작업 내역
    ws = wb.create_sheet("전체 작업 내역")
    header(ws, 1, ["방송사", "장르", "프로그램명", "회차", "작업구분", "담당자", "단가", "영상 길이(분)",
                   "정산 금액", "작업소요(분)", "최초 완료(KST)", "prj_id", "비고"],
           [8, 7, 30, 14, 8, 9, 10, 11, 12, 11, 16, 8, 40])
    ws.freeze_panes = "A2"
    fmts = [None] * 6 + [WON, "#,##0", WON, "0.0", None, None, None]
    for i, x in enumerate(rows, 2):
        put(ws, i, [x["brd"], x["genre"], x["program"], x["ep"], x["work"], x["assignee"],
                    f"=SUMIFS(단가표!$C$4:$C$15,단가표!$A$4:$A$15,A{i},단가표!$B$4:$B$15,E{i})",
                    x["dur"], f"=G{i}*H{i}", x["work_min"], x["done"], x["prj_id"], x["note"]],
            fmts, WARN_FILL if x["warn"] else None)
        x["row"] = i
    ws.auto_filter.ref = f"A1:M{last}"
    t = last + 1
    put(ws, t, ["합계", None, None, None, f'=COUNTA(E2:E{last})&"건"', None, None,
                f"=SUM(H2:H{last})", f"=SUM(I2:I{last})", f"=SUM(J2:J{last})", None, None, None],
        fmts)
    for c in ws[t]:
        c.font = Font(bold=True)

    # 방송사별
    ws = wb.create_sheet("방송사별")
    ws["A1"] = "방송사별 작업 내역"; ws["A1"].font = Font(bold=True, size=13)
    header(ws, 3, ["방송사", "총 건수", "싱크 건수", "교정 건수", "누적 영상 시간(분)", "누적 정산 금액", "작업자 수"],
           [10, 9, 10, 10, 16, 16, 10])
    fm = ["#,##0"] * 5 + [WON, "#,##0"]
    for i, b in enumerate(ORDER, 4):
        people = {x["assignee"] for x in rows if x["brd"] == b and x["assignee"]}
        put(ws, i, [b, f"=COUNTIFS({A},A{i})", f'=COUNTIFS({A},A{i},{E},"싱크")',
                    f'=COUNTIFS({A},A{i},{E},"교정")', f"=SUMIFS({col('H')},{A},A{i})",
                    f"=SUMIFS({col('I')},{A},A{i})", len(people)], fm)
    s = 4 + len(ORDER)
    put(ws, s, ["합계"] + [f"=SUM({L}4:{L}{s - 1})" for L in "BCDEF"]
        + [len({x["assignee"] for x in rows if x["assignee"]})], fm)
    for c in ws[s]:
        c.font = Font(bold=True)
    ws.cell(s + 2, 1, f"작업자 수 = 해당 방송사에서 {month} 정산 건이 있는 담당자 수(중복 제외). 합계 행은 전체 고유 인원.")

    # 작업자별 요약
    people = sorted({x["assignee"] for x in rows if x["assignee"]})
    ws = wb.create_sheet("작업자별 요약")
    ws["A1"] = "작업자별 요약"; ws["A1"].font = Font(bold=True, size=13)
    header(ws, 3, ["담당자", "총 건수", "싱크 건수", "교정 건수", "누적 영상 시간(분)", "누적 정산 금액", "작업소요(분)"],
           [12, 9, 10, 10, 16, 16, 13])
    fm = ["#,##0"] * 5 + [WON, "#,##0.0"]
    for i, p in enumerate(people, 4):
        put(ws, i, [p, f"=COUNTIFS({F},A{i})", f'=COUNTIFS({F},A{i},{E},"싱크")',
                    f'=COUNTIFS({F},A{i},{E},"교정")', f"=SUMIFS({col('H')},{F},A{i})",
                    f"=SUMIFS({col('I')},{F},A{i})", f"=SUMIFS({col('J')},{F},A{i})"], fm)
    s = 4 + len(people)
    put(ws, s, ["합계"] + [f"=SUM({L}4:{L}{s - 1})" for L in "BCDEFG"], fm)
    for c in ws[s]:
        c.font = Font(bold=True)

    # 작업자별 시트
    for p in people:
        mine = [x for x in rows if x["assignee"] == p]
        ws = wb.create_sheet(p[:31])
        ws["A1"] = f"{p} — {month} 작업 내역"; ws["A1"].font = Font(bold=True, size=13)
        header(ws, 3, ["방송사", "장르", "프로그램명", "회차", "작업구분", "단가", "영상 길이(분)", "정산 금액",
                       "작업소요(분)", "최초 완료(KST)"], [10, 7, 30, 14, 8, 10, 11, 12, 11, 16])
        ws.freeze_panes = "A4"
        fm = [None] * 5 + [WON, None, WON, "0.0", None]
        for i, x in enumerate(mine, 4):
            put(ws, i, [x["brd"], x["genre"], x["program"], x["ep"], x["work"],
                        f"='전체 작업 내역'!G{x['row']}", x["dur"], f"='전체 작업 내역'!I{x['row']}",
                        x["work_min"], x["done"]], fm, WARN_FILL if x["warn"] else None)
        end = 3 + len(mine)
        r = end + 2
        ws.cell(r, 1, "방송사별 총 건수 및 금액").font = Font(bold=True)
        header(ws, r + 1, ["방송사", "건수", "영상 길이(분)", "정산 금액"])
        rng = lambda L: f"${L}$4:${L}${end}"
        fm2 = [None, "#,##0", "#,##0", WON]
        brds = [b for b in ORDER if any(x["brd"] == b for x in mine)]
        for k, b in enumerate(brds, r + 2):
            put(ws, k, [b, f"=COUNTIFS({rng('A')},A{k})", f"=SUMIFS({rng('G')},{rng('A')},A{k})",
                        f"=SUMIFS({rng('H')},{rng('A')},A{k})"], fm2)
        k = r + 2 + len(brds)
        put(ws, k, ["전체", f"=COUNTA({rng('E')})", f"=SUM({rng('G')})", f"=SUM({rng('H')})"], fm2)
        for c in ws[k][:4]:
            c.font = Font(bold=True)

    wb.save(out)
    return n


if __name__ == "__main__":
    src, out = sys.argv[1], sys.argv[2]
    rows, skipped = load(src)
    print(build(rows, skipped, out, sys.argv[3] if len(sys.argv) > 3 else "9월"), "rows,", len(skipped), "skipped")
