#!/usr/bin/env python3
"""월 정산 워크북 — 정산 쿼리 결과를 '(수정중)정산_워크북_틀' 형식으로 채운다.

사용: python3 settlement_workbook.py <틀.xlsx> <정산.tsv|psv> <YYYY-MM> <출력.xlsx>
- 데이터는 정산_쿼리.sh 출력(탭) 또는 붙여넣기를 옮긴 | 구분. 교정 최초 완료가 YYYY-MM 인 회차만 쓴다.
- 틀의 '전체 작업 리스트' 행은 비우고 해당 월 회차로 채운다(회차당 1행, 교정 완료일순).
- 틀 행에 묶인 수동 단가 수식(191행 등)은 표준 수식으로 되돌린다.
- 프로그램명은 틀(상세 시트 목록·기존 행·별칭)에 맞춰 통일, 빈 장르는 틀에서 보충.
- 'OOO 관리자/작업자' 내부 계정이 담당자면 완료처리자가 실제 작업자일 때 그 사람으로 바꾸고 비고에 남김.
- 구글 시트 전용 수식(__xludf.DUMMYFUNCTION)의 캐시값을 계산해 넣음 → 엑셀에서도 값이 보임.
"""
import copy
import csv
import re
import sys
from collections import Counter
from datetime import datetime

import openpyxl

BRD = {"딜라이브": "DLIV", "KBS": "KBS", "LG헬로비전": "LGHV", "SK브로드밴드": "SKBB",
       "TVCS": "TVCS", "티캐스트": "TVCS", "TVING": "TVNG"}
# DB 등록명 → 틀 표기 (9월 정산 때 사람이 맞춘 것 + 별칭표)
NAME = {"경기도의회의정토크쇼": "의정토크쇼", "눈에띄는그녀들9": "눈에 띄는 그녀들",
        "당신의 골목 맛집 당골집": "당신의 골목 맛집, 당골집", "당신의 골목 맛집, 당골집2": "당신의 골목 맛집, 당골집",
        "전현무 계획3": "전현무 계획", "전현무계획": "전현무 계획", "태군 노래자랑 경북 경산": "태군노래자랑",
        "i-로드 인천문화기행": "아이로드", "주말여행 산이 좋다 시즌3": "주말여행 산이좋다"}
GENRE_FIX = {"기쁜 우리 좋은날": "드라마"}   # t_prg 는 예능으로 잘못 등록, 회차 장르는 드라마
LISTS = {"DLIV": ("H", "I"), "KBS": ("M", "N"), "LGHV": ("R", "S"), "SKBB": ("W", "X"),
         "TVCS": ("AB", "AC"), "TVNG": ("AG", "AH")}   # 상세 시트 방송사별 프로그램 목록 열
MASTER_ROWS = range(2, 254)   # 틀 '전체 작업 리스트' 데이터 행
COLS = "ABCDEIJLNQ"           # 입력 열 (나머지는 틀 수식)

PAT = re.compile(r"^(=IFERROR\(__xludf\.DUMMYFUNCTION\(.*\)),(.*)\)$", re.S)
internal = lambda n: not n or n.endswith((" 관리자", " 작업자"))   # noqa: E731


# ── 구글 시트 전용 수식 캐시 ──────────────────────────────────────

def _d(v):
    return v.strftime("%Y-%m-%d") if isinstance(v, datetime) else str(v)[:10]


def rate_rows(wb):
    return [r for r in wb["단가표"].iter_rows(min_row=2, max_row=21, values_only=True) if r[0]]


def money(rates, a, b, c, i, length, col):
    """틀 K/M/O 수식과 같은 계산. col 4=싱크 5=교정 6=BF. (틀 캐시 2,989개와 일치 검증)"""
    if not a:
        return ""
    w = _d(i)
    hit = [r for r in rates if r[0] == a and r[1] == b and (r[2] == c or not r[2]) and _d(r[8]) <= w <= _d(r[9])]
    hit.sort(key=lambda r: (bool(r[2]), _d(r[8])), reverse=True)
    rate = hit[0][col] if hit else ""
    if length in ("", None):
        return ""
    if length == 0:
        return 0
    if rate in ("", None) or str(rate).strip() == "-" or str(length).strip() == "-":
        return "확인필요"
    return 0 if rate == 0 else rate * length


def lit(v):
    return repr(float(v)) if isinstance(v, (int, float)) else '"' + str(v).replace('"', '""') + '"'


def set_cache(cell, value):
    m = PAT.match(cell.value)
    assert m, cell.coordinate
    cell.value = f"{m.group(1)},{lit(value)})"


# ── 데이터 ───────────────────────────────────────────────────────

def load(path, ym, canon, tmpl_genre):
    with open(path, encoding="utf-8-sig") as f:
        lines = [ln for ln in f if ln.strip() and not ln.startswith("#")]
    rows = list(csv.DictReader(lines, delimiter="\t" if "\t" in lines[0] else "|"))
    by = {}
    for r in rows:
        by.setdefault(r["prj_id"], []).append(r)
    eps, skipped = [], []
    for pid, rs in by.items():
        e = next((r for r in rs if r["work"] == "EDIT"), None)
        if not e or not e["first_done_kst"].startswith(ym):
            continue
        b = BRD.get(e["brd"])
        if not b or re.search(r"TEST|테스트", e["program"], re.I):
            skipped.append(f"{e['brd']} {e['program']} {e['episode']}회 (prj {pid})")
            continue
        s = next((r for r in rs if r["work"] == "SYNC"), None)
        notes = ["장르별(분리작업)"] if b in ("SKBB", "LGHV") else []

        def who(r, label):
            if r is None:
                return None
            a, done = r["assignee"], r["done_by"]
            if internal(a) and not internal(done):
                notes.append(f"{label} 담당 {a or '미배정'} → 완료처리자 {done}")
                return done
            if internal(a):
                notes.append(f"{label} 담당 {a or done} — 내부 계정, 실제 작업자 확인 필요")
                return a or done
            return a

        c = canon(b, e["program"])
        lab = (f"시즌{e['season']} " if e["season"] else "") + (f"{e['episode']}회" if e["episode"] else "")
        eps.append(dict(A=b, B=GENRE_FIX.get(c) or e["genre"] or tmpl_genre.get((b, c)), C=c, D=lab or None,
                        E=int(e["dur_min_ceil"]) if e["dur_min_ceil"] else None,
                        I=datetime.strptime(e["first_done_kst"], "%Y-%m-%d %H:%M"),
                        J=who(s, "싱크"), L=who(e, "교정"), N=None, pid=pid, notes=notes))
    dup = Counter((x["A"], x["C"], x["D"]) for x in eps)
    for x in eps:
        if dup[(x["A"], x["C"], x["D"])] > 1:
            x["notes"].append(f"같은 회차 프로젝트 중복(prj {x['pid']}) 확인 필요")
        x["Q"] = " / ".join(x["notes"]) or None
    eps.sort(key=lambda x: x["I"])
    return eps, skipped


def main(src, data, ym, out):
    wb = openpyxl.load_workbook(src)
    ws, d = wb["전체 작업 리스트"], wb["방송사별 상세(확인전)"]

    canon_set = {b: [d[f"{col}{r}"].value for r in range(40, 99) if d[f"{col}{r}"].value] for b, (_, col) in LISTS.items()}
    tmpl_genre = {}
    for i in MASTER_ROWS:
        a, g, c = ws[f"A{i}"].value, ws[f"B{i}"].value, ws[f"C{i}"].value
        if a and c:
            tmpl_genre.setdefault((a, c), g)
            if c not in canon_set.setdefault(a, []):
                canon_set[a].append(c)

    def canon(b, name):
        name = NAME.get(name, name)
        ns = lambda t: re.sub(r"\s", "", t)   # noqa: E731
        for c in canon_set.get(b, []):
            if ns(c) == ns(name):
                return c
        pre = [c for c in canon_set.get(b, []) if ns(name).startswith(ns(c))]   # '1+1다큐…(10) 부제' → 본 제목
        return max(pre, key=len) if pre else name

    eps, skipped = load(data, ym, canon, tmpl_genre)
    assert len(eps) < len(MASTER_ROWS), "틀 행 수(252)를 넘음 — 틀 수식 범위부터 늘려야 함"
    print(f"{ym} 회차 {len(eps)} / 제외 {skipped}")

    # 전체 작업 리스트
    for i in MASTER_ROWS:
        for c in COLS:
            ws[f"{c}{i}"].value = None
    for i, x in enumerate(eps, 2):
        for c in COLS:
            ws[f"{c}{i}"].value = x[c]
    for i in MASTER_ROWS:   # 행에 묶인 수동 단가 수식 → 바로 윗행 표준 수식에서 행 번호만 바꿔 복원
        for col in "KM":
            if not PAT.match(ws[f"{col}{i}"].value or ""):
                body = PAT.match(ws[f"{col}{i - 1}"].value).group(1)
                ws[f"{col}{i}"].value = re.sub(rf"(?<![0-9$]){i - 1}(?![0-9])", str(i), body) + ',"")'
                print(f"수동 단가 수식 복원: {col}{i}")
    rates = rate_rows(wb)
    for i in MASTER_ROWS:
        a, b, c, e, it = (ws[f"{k}{i}"].value for k in "ABCEI")
        for col, w, idx in (("K", "J", 4), ("M", "L", 5), ("O", "N", 6)):
            set_cache(ws[f"{col}{i}"], money(rates, a, b, c, it, e if ws[f"{w}{i}"].value else "", idx))

    # 제작원별 요약 — 명단에 없는 실제 작업자 추가 (내부 계정 제외)
    p = wb["제작원별 요약"]
    tot = next(r for r in range(4, 60) if p[f"B{r}"].value == "합계")
    names = {p[f"B{r}"].value for r in range(4, tot)}
    new_people = sorted({x[k] for x in eps for k in "JLN" if x[k] and not internal(x[k])} - names)
    for person in new_people:
        for col in "BCDEFGHI":
            p[f"{col}{tot + 1}"]._style = copy.copy(p[f"{col}{tot}"]._style)
            p[f"{col}{tot + 1}"].value = "합계" if col == "B" else f"=SUM({col}4:{col}{tot})"
        for col in "BCDEFGHI":
            v = p[f"{col}{tot - 1}"].value
            if isinstance(v, str):
                v = v.replace(f"$B{tot - 1}", f"$B{tot}").replace(f"C{tot - 1}:H{tot - 1}", f"C{tot}:H{tot}")
            p[f"{col}{tot}"].value = v
            p[f"{col}{tot}"]._style = copy.copy(p[f"{col}{tot - 1}"]._style)
        p[f"B{tot}"].value = person
        tot += 1
    print("제작원별 요약 추가", new_people)

    # 방송사별 요약 작업자 수 캐시
    s = wb["방송사별 요약"]
    master = [{k: ws[f"{k}{i}"].value for k in "AJLN"} for i in MASTER_ROWS]
    for r in range(4, 10):
        brd = s[f"B{r}"].value
        set_cache(s[f"F{r}"], len({x[k] for x in master for k in "JLN" if x[k] and (brd == "합계" or x["A"] == brd)}))

    # 상세 시트 — 새 프로그램을 방송사 목록에 추가, 현재 선택(C2) 목록 캐시
    for x in eps:
        hc, col = LISTS[x["A"]]
        cur = [d[f"{col}{r}"].value for r in range(40, 99)]
        if x["C"] not in cur:
            r = 40 + cur.index(None)
            d[f"{hc}{r}"].value, d[f"{col}{r}"].value = x["A"], x["C"]
            print("상세 목록 추가", x["A"], x["C"])
    sel_col = LISTS[d["C2"].value][1]
    lst = [d[f"{sel_col}{r}"].value for r in range(40, 99) if d[f"{sel_col}{r}"].value]
    for k, r in enumerate(range(6, 999)):
        if isinstance(d[f"B{r}"].value, str) and "DUMMYFUNCTION" in d[f"B{r}"].value:
            set_cache(d[f"B{r}"], lst[k] if k < len(lst) else "")

    # 제작원별_틀 — 정산 월·기간·작업자 목록
    t = wb["제작원별_틀"]
    y, m = map(int, ym.split("-"))
    t["G3"].value, t["I3"].value = f"{m:02d}", y
    last_day = (datetime(y + (m == 12), m % 12 + 1, 1) - datetime(y, m, 1)).days
    set_cache(t["C4"], f"{ym}-01 ~ {ym}-{last_day:02d}")
    people = sorted({x[k] for x in master for k in "JLN" if x[k]}
                    | {v for v in (p[f"B{r}"].value for r in range(4, 14)) if v and v != "합계"})
    for k in range(max(len(people), 20)):
        cell, v = t[f"N{7 + k}"], people[k] if k < len(people) else ""
        if k == 0:
            set_cache(cell, v)
        elif v or cell.value:
            cell.value = f'=IFERROR(__xludf.DUMMYFUNCTION("""COMPUTED_VALUE"""),{lit(v)})'
    wb.save(out)
    print("저장", out)


if __name__ == "__main__":
    if len(sys.argv) != 5:
        sys.exit(__doc__)
    main(*sys.argv[1:])
